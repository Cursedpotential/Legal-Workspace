"""Whole office work products with immutable bytes and transactional revisions.

Uses the workspace's existing legal.sqlite and audit outbox. These are private
work products, never imported evidence or released artifacts.
"""

from __future__ import annotations

import hashlib
import io
import json
import sqlite3
import time
import zipfile
from contextlib import contextmanager
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

from docx import Document
from pydantic import BaseModel

from legal_workspace.db.store import create_tables

MAX_DOCUMENT_BYTES = 32 * 1024 * 1024
MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "odt": "application/vnd.oasis.opendocument.text",
}


class DocumentNotFound(KeyError):
    pass


class DocumentLocked(ValueError):
    pass


class RevisionConflict(ValueError):
    def __init__(self, current_revision: int):
        self.current_revision = current_revision
        super().__init__("The document has changed. Reload before saving.")


class WorkDocument(BaseModel):
    document_id: str
    matter_id: str
    title: str
    filename: str
    format: str
    current_revision: int
    original_revision: int = 1
    created_at: str
    updated_at: str


class WorkDocumentRevision(BaseModel):
    document_id: str
    revision: int
    parent_revision: int | None
    sha256: str
    byte_length: int
    actor: str
    origin: str
    created_at: str


def validate_office_bytes(data: bytes, format: str) -> None:
    if format not in MEDIA_TYPES:
        raise ValueError("Only DOCX and ODT work products are supported.")
    if not data or len(data) > MAX_DOCUMENT_BYTES:
        raise ValueError("Document must be between 1 byte and 32 MiB.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 10000 or sum(item.file_size for item in entries) > 256 * 1024 * 1024:
                raise ValueError("Office archive is too large when expanded.")
            names = set(archive.namelist())
            required = (
                {"[Content_Types].xml", "word/document.xml"}
                if format == "docx"
                else {"mimetype", "content.xml"}
            )
            if not required <= names:
                raise ValueError("File contents do not match the office format.")
            if format == "odt" and archive.read("mimetype") != MEDIA_TYPES["odt"].encode():
                raise ValueError("File is not an ODT text document.")
    except zipfile.BadZipFile as exc:
        raise ValueError("Invalid office document archive.") from exc


class WorkDocumentService:
    def __init__(self, store_dir: Path, matter_id: str):
        self.store_dir = Path(store_dir)
        self.matter_id = str(matter_id)
        create_tables(store_dir=str(self.store_dir))
        self.db_path = self.store_dir / "legal.sqlite"
        with self._connect() as conn:
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS legal_work_document (
                    document_id TEXT PRIMARY KEY, matter_id TEXT NOT NULL,
                    title TEXT NOT NULL, filename TEXT NOT NULL, format TEXT NOT NULL,
                    current_revision INTEGER NOT NULL, original_revision INTEGER NOT NULL,
                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS legal_work_document_revision (
                    document_id TEXT NOT NULL REFERENCES legal_work_document(document_id),
                    revision INTEGER NOT NULL, parent_revision INTEGER,
                    sha256 TEXT NOT NULL, byte_length INTEGER NOT NULL,
                    actor TEXT NOT NULL, origin TEXT NOT NULL, created_at TEXT NOT NULL,
                    content BLOB NOT NULL, PRIMARY KEY(document_id, revision));
            """)

    @contextmanager
    def _connect(self):
        conn = sqlite3.connect(self.db_path, timeout=30)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys=ON")
        conn.execute("PRAGMA synchronous=FULL")
        try:
            with conn:
                yield conn
        finally:
            conn.close()

    def _get(self, conn, document_id: str) -> WorkDocument:
        row = conn.execute(
            "SELECT * FROM legal_work_document WHERE document_id=? AND matter_id=?",
            (document_id, self.matter_id),
        ).fetchone()
        if row is None:
            raise DocumentNotFound(document_id)
        return WorkDocument(**dict(row))

    def get(self, document_id: str) -> WorkDocument:
        with self._connect() as conn:
            return self._get(conn, document_id)

    def list(self) -> list[WorkDocument]:
        with self._connect() as conn:
            return [
                WorkDocument(**dict(row))
                for row in conn.execute(
                    "SELECT * FROM legal_work_document WHERE matter_id=? ORDER BY updated_at DESC, document_id",
                    (self.matter_id,),
                )
            ]

    def create_blank(self, title: str, *, actor: str) -> WorkDocument:
        buffer = io.BytesIO()
        Document().save(buffer)
        return self.import_document(
            "Untitled.docx", buffer.getvalue(), title=title, actor=actor, origin="blank"
        )

    def import_document(
        self,
        filename: str,
        data: bytes,
        *,
        title: str | None = None,
        actor: str,
        origin: str = "import",
    ) -> WorkDocument:
        filename = filename.replace("\\", "/").split("/")[-1]
        if not filename or any(ord(char) < 32 for char in filename):
            raise ValueError("A valid filename is required.")
        format = Path(filename).suffix.lower().lstrip(".")
        validate_office_bytes(data, format)
        title = (title or Path(filename).stem).strip()
        if not title or len(title) > 250:
            raise ValueError("Title must contain 1 to 250 characters.")
        stamp = datetime.now(UTC).isoformat()
        document = WorkDocument(
            document_id=str(uuid4()),
            matter_id=self.matter_id,
            title=title,
            filename=filename,
            format=format,
            current_revision=1,
            created_at=stamp,
            updated_at=stamp,
        )
        with self._connect() as conn:
            conn.execute("BEGIN IMMEDIATE")
            conn.execute(
                "INSERT INTO legal_work_document VALUES (?,?,?,?,?,?,?,?,?)",
                tuple(document.model_dump().values()),
            )
            self._revision(conn, document, data, actor, origin, None)
        return document

    def _revision(
        self, conn, document: WorkDocument, data: bytes, actor: str, origin: str, parent: int | None
    ):
        revision = WorkDocumentRevision(
            document_id=document.document_id,
            revision=document.current_revision,
            parent_revision=parent,
            sha256=hashlib.sha256(data).hexdigest(),
            byte_length=len(data),
            actor=actor,
            origin=origin,
            created_at=document.updated_at,
        )
        conn.execute(
            "INSERT INTO legal_work_document_revision VALUES (?,?,?,?,?,?,?,?,?)",
            (*revision.model_dump().values(), data),
        )
        payload = json.dumps(revision.model_dump(), sort_keys=True)
        conn.execute(
            """INSERT INTO legal_audit_event_outbox
            (event_id,event_type,schema_version,occurred_at,matter_id,court_case_id,
             aggregate_id,aggregate_version,trace_id,payload_hash,payload,published_at)
            VALUES (?,?,'1',?,?,NULL,?,?,?,?,?,NULL)""",
            (
                str(uuid4()),
                "legal.work_document.revision_created",
                document.updated_at,
                self.matter_id,
                document.document_id,
                document.current_revision,
                str(uuid4()),
                "sha256:" + hashlib.sha256(payload.encode()).hexdigest(),
                payload,
            ),
        )

    def read_bytes(self, document_id: str, revision: int | None = None) -> bytes:
        with self._connect() as conn:
            document = self._get(conn, document_id)
            row = conn.execute(
                "SELECT content FROM legal_work_document_revision WHERE document_id=? AND revision=?",
                (document_id, revision if revision is not None else document.current_revision),
            ).fetchone()
            if row is None:
                raise DocumentNotFound(f"{document_id} revision {revision}")
            return bytes(row["content"])

    def history(self, document_id: str) -> list[WorkDocumentRevision]:
        with self._connect() as conn:
            self._get(conn, document_id)
            return [
                WorkDocumentRevision(**dict(row))
                for row in conn.execute(
                    "SELECT document_id,revision,parent_revision,sha256,byte_length,actor,origin,created_at FROM legal_work_document_revision WHERE document_id=? ORDER BY revision DESC",
                    (document_id,),
                )
            ]

    def save(
        self,
        document_id: str,
        data: bytes,
        expected_revision: int,
        *,
        actor: str,
        origin: str = "office",
        office_lock: str | None = None,
    ) -> WorkDocument:
        with self._connect() as conn:
            conn.execute("BEGIN IMMEDIATE")
            document = self._get(conn, document_id)
            lock = None
            if conn.execute(
                "SELECT 1 FROM sqlite_master WHERE type='table' AND name='office_editor_locks'"
            ).fetchone():
                lock = conn.execute(
                    "SELECT * FROM office_editor_locks WHERE document_id=? AND expires>?",
                    (document_id, time.time()),
                ).fetchone()
                if lock and lock["lock_value"] != office_lock:
                    raise DocumentLocked("Document is locked by the office editor.")
            if office_lock and lock is None:
                raise DocumentLocked("Office editor lock has expired. Reconnect before saving.")
            if document.current_revision != expected_revision:
                raise RevisionConflict(document.current_revision)
            validate_office_bytes(data, document.format)
            document.current_revision += 1
            document.updated_at = datetime.now(UTC).isoformat()
            conn.execute(
                "UPDATE legal_work_document SET current_revision=?,updated_at=? WHERE document_id=?",
                (document.current_revision, document.updated_at, document_id),
            )
            self._revision(conn, document, data, actor, origin, expected_revision)
            if lock and office_lock:
                conn.execute(
                    "UPDATE office_editor_locks SET revision=? WHERE document_id=?",
                    (document.current_revision, document_id),
                )
        return document
