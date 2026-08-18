"""Write a local DOCX for a release candidate. Not a filing.

> _Byline: Grok · grok-4.6 · 2026-08-18_
PDF waits for the renderer sidecar. This is the Type 2 local stand-in.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from legal_workspace.domain.release import ReleaseManifest


def write_release_docx(
    path: Path,
    manifest: ReleaseManifest,
    sections: list[object],
) -> Path:
    document = Document()
    document.add_heading("Release candidate", level=1)
    document.add_paragraph("This is not a filing and not court-safe.")
    document.add_paragraph(f"content_hash: {manifest.content_hash}")
    document.add_paragraph(f"package_id: {manifest.package_id}")
    document.add_paragraph(f"state: {manifest.state.value}")
    document.add_paragraph("filed: false")
    document.add_paragraph("Omitted: " + ", ".join(manifest.omitted_private))
    for section in sections:
        document.add_heading(str(getattr(section, "heading")), level=2)
        document.add_paragraph(str(getattr(section, "body")))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path
