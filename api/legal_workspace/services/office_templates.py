"""Versioned DOCX drafting templates built from the reviewed outline catalog."""

from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass
from typing import Mapping, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from legal_workspace.domain.templates import DraftingTemplate, get_template

TEMPLATE_VERSION = "1.0.0"


@dataclass(frozen=True)
class OfficeTemplateInfo:
    template_id: str
    version: str
    title: str
    description: str
    kind: str
    filename: str
    fields: tuple[str, ...]


@dataclass(frozen=True)
class OfficeTemplateArtifact:
    info: OfficeTemplateInfo
    content: bytes
    content_hash: str


def _fields(template: DraftingTemplate) -> tuple[str, ...]:
    return (
        "caption.case_name", "caption.case_number", "caption.court",
        "caption.judge", "sections", "signature.name", "signature.address",
        "signature.phone", "signature.email",
    )


def template_info(template_id: str) -> OfficeTemplateInfo:
    template = get_template(template_id)
    return OfficeTemplateInfo(
        template_id=template.template_id, version=TEMPLATE_VERSION,
        title=template.title, kind=template.kind.value,
        description=f"Editable {template.kind.value.replace('_', ' ')} working document with caption, section, and signature fields.",
        filename=f"{template.template_id}-v{TEMPLATE_VERSION}.docx",
        fields=_fields(template),
    )


def template_catalog() -> list[OfficeTemplateInfo]:
    from legal_workspace.domain.templates import catalog
    return [template_info(item.template_id) for item in catalog()]


def _value(values: Mapping[str, object], key: str, default: str) -> str:
    raw = values.get(key)
    return str(raw).strip() if raw is not None and str(raw).strip() else default


def _caption_table(document: Document, caption: Mapping[str, object]) -> None:
    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    labels = (
        ("Case name", _value(caption, "case_name", "[Enter case name]")),
        ("Case number", _value(caption, "case_number", "[Enter case number]")),
        ("Court", _value(caption, "court", "[Enter court]")),
        ("Judge", _value(caption, "judge", "[Enter judge]")),
    )
    for row, (label, value) in zip(table.rows, labels):
        row.cells[0].text = label
        row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].bold = True


def _signature_block(document: Document, signature: Mapping[str, object]) -> None:
    document.add_heading("Signature", level=2)
    for label, key in (("Name", "name"), ("Address", "address"),
                       ("Phone", "phone"), ("Email", "email")):
        document.add_paragraph(
            f"{label}: {_value(signature, key, f'[Enter {label.lower()}]')}"
        )
    document.add_paragraph("Signature: ____________________________________")
    document.add_paragraph("Date: _________________________________________")


def build_template_document(
    template_id: str,
    *,
    title: str | None = None,
    caption: Mapping[str, object] | None = None,
    sections: Sequence[str] | None = None,
    signature: Mapping[str, object] | None = None,
) -> OfficeTemplateArtifact:
    """Build a styled DOCX retaining every outline section for editing."""
    template = get_template(template_id)
    info = template_info(template_id)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    heading = document.add_heading(title or template.title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"Template version: {info.version}")
    document.add_paragraph("Working document. Complete, review, and revise each field before use.")
    document.add_heading("Caption", level=2)
    _caption_table(document, caption or {})
    document.add_heading("Sections", level=2)
    supplied = list(sections or ())
    for index, outline in enumerate(template.outline):
        heading = document.add_heading(outline.split(":", 1)[0], level=3)
        body = supplied[index].strip() if index < len(supplied) and supplied[index].strip() else f"[Enter content for: {outline}]"
        document.add_paragraph(body)
        heading.paragraph_format.keep_with_next = True
    _signature_block(document, signature or {})

    stream = io.BytesIO()
    document.save(stream)
    content = stream.getvalue()
    return OfficeTemplateArtifact(
        info=info, content=content,
        content_hash=hashlib.sha256(content).hexdigest(),
    )
