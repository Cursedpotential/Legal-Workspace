import sqlite3
import time
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile

import pytest
from docx import Document
from legal_workspace.services.work_documents import (
    DocumentLocked,
    DocumentNotFound,
    RevisionConflict,
    WorkDocumentService,
)


def docx_bytes(text="Original"):
    buffer = BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def service(tmp_path):
    return WorkDocumentService(tmp_path, str(uuid4()))


def test_blank_import_revisions_survive_recreation(service):
    blank = service.create_blank("Motion", actor="owner")
    Document(BytesIO(service.read_bytes(blank.document_id)))
    original = docx_bytes()
    imported = service.import_document("Draft.docx", original, actor="owner")
    changed = docx_bytes("Revised")
    saved = service.save(imported.document_id, changed, 1, actor="owner")
    recreated = WorkDocumentService(service.store_dir, service.matter_id)
    assert saved.current_revision == 2
    assert len(recreated.list()) == 2
    assert recreated.read_bytes(imported.document_id) == changed
    assert recreated.read_bytes(imported.document_id, 1) == original
    assert [row.parent_revision for row in recreated.history(imported.document_id)] == [1, None]
    with sqlite3.connect(service.db_path) as conn:
        assert (
            conn.execute(
                "SELECT count(*) FROM legal_audit_event_outbox WHERE event_type='legal.work_document.revision_created'"
            ).fetchone()[0]
            == 3
        )


def test_odt_import_preserves_original(service):
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr(
            "content.xml",
            '<office:document xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"/>',
        )
    data = buffer.getvalue()
    document = service.import_document("notes.odt", data, actor="owner")
    assert service.read_bytes(document.document_id) == data
    assert document.format == "odt"


def test_same_filename_creates_distinct_documents_and_missing_revision_is_explicit(service):
    first = service.import_document("draft.docx", docx_bytes("First"), actor="owner")
    second = service.import_document("draft.docx", docx_bytes("Second"), actor="owner")
    assert first.document_id != second.document_id
    assert service.read_bytes(first.document_id) != service.read_bytes(second.document_id)
    with pytest.raises(DocumentNotFound):
        service.read_bytes(first.document_id, 99)


def test_stale_and_concurrent_saves_cannot_overwrite(service):
    document = service.create_blank("Concurrent", actor="owner")
    data = docx_bytes("Changed")

    def save(_):
        other = WorkDocumentService(service.store_dir, service.matter_id)
        try:
            return other.save(document.document_id, data, 1, actor="owner").current_revision
        except RevisionConflict as exc:
            return f"conflict:{exc.current_revision}"

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(save, range(2)))
    assert sorted(results, key=str) == [2, "conflict:2"]
    assert len(service.history(document.document_id)) == 2


def test_invalid_import_save_and_matter_boundary(service):
    with pytest.raises(ValueError):
        service.import_document("fake.docx", b"bad", actor="owner")
    document = service.create_blank("Valid", actor="owner")
    with pytest.raises(ValueError):
        service.save(document.document_id, b"bad", 1, actor="owner")
    assert service.get(document.document_id).current_revision == 1
    other = WorkDocumentService(service.store_dir, str(uuid4()))
    assert other.list() == []
    with pytest.raises(DocumentNotFound):
        other.read_bytes(document.document_id)


def test_audit_failure_rolls_back_bytes_and_revision(service):
    document = service.create_blank("Atomic", actor="owner")
    with sqlite3.connect(service.db_path) as conn:
        conn.execute(
            "CREATE TRIGGER reject_document_audit BEFORE INSERT ON legal_audit_event_outbox BEGIN SELECT RAISE(ABORT, 'audit unavailable'); END"
        )
    with pytest.raises(sqlite3.IntegrityError):
        service.save(document.document_id, docx_bytes("Unsaved"), 1, actor="owner")
    assert service.get(document.document_id).current_revision == 1
    assert len(service.history(document.document_id)) == 1


def test_office_lock_blocks_upload_and_advances_with_save(service):
    document = service.create_blank("Locked", actor="owner")
    with sqlite3.connect(service.db_path) as conn:
        conn.execute(
            "CREATE TABLE office_editor_locks(document_id TEXT PRIMARY KEY, lock_value TEXT NOT NULL, expires REAL NOT NULL, revision INTEGER NOT NULL)"
        )
        conn.execute(
            "INSERT INTO office_editor_locks VALUES (?,?,?,?)",
            (document.document_id, "lock", time.time() + 60, 1),
        )
    data = docx_bytes("Office change")
    with pytest.raises(DocumentLocked):
        service.save(document.document_id, data, 1, actor="owner")
    service.save(document.document_id, data, 1, actor="owner", office_lock="lock")
    with sqlite3.connect(service.db_path) as conn:
        assert conn.execute("SELECT revision FROM office_editor_locks").fetchone()[0] == 2
        conn.execute("UPDATE office_editor_locks SET expires=?", (time.time() - 60,))
    with pytest.raises(DocumentLocked):
        service.save(document.document_id, data, 2, actor="owner", office_lock="lock")
    assert service.get(document.document_id).current_revision == 2
    assert service.save(document.document_id, data, 2, actor="owner").current_revision == 3


def test_http_auth_download_and_conflict(service):
    from fastapi import FastAPI, Request
    from fastapi.testclient import TestClient
    from legal_workspace.api.auth import AuthenticatedPrincipal
    from legal_workspace.api.work_document_routes import get_document_service, router

    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_document_service] = lambda: service
    unauthenticated = FastAPI()
    unauthenticated.include_router(router)
    unauthenticated.dependency_overrides[get_document_service] = lambda: service
    with TestClient(unauthenticated) as client:
        assert client.get("/v1/work-documents").status_code == 401

    @app.middleware("http")
    async def authenticate(request: Request, call_next):
        request.state.auth = AuthenticatedPrincipal("human", None, None, (), "test")
        return await call_next(request)

    with TestClient(app) as client:
        created = client.post("/v1/work-documents", json={"title": "HTTP document"})
        assert created.status_code == 201
        document_id = created.json()["document_id"]
        response = client.get(f"/v1/work-documents/{document_id}/content")
        assert response.status_code == 200
        Document(BytesIO(response.content))
        args = {
            "files": {"file": ("updated.docx", docx_bytes())},
            "data": {"expected_revision": "1"},
        }
        assert client.put(f"/v1/work-documents/{document_id}/content", **args).status_code == 200
        conflict = client.put(f"/v1/work-documents/{document_id}/content", **args)
        assert conflict.status_code == 409
        assert conflict.json()["detail"]["current_revision"] == 2
