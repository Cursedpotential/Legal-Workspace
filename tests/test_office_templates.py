from io import BytesIO
import hashlib

from docx import Document

from legal_workspace.services.office_templates import (
    TEMPLATE_VERSION, build_template_document, template_catalog,
)


def test_catalog_is_versioned_and_exposes_editable_fields() -> None:
    rows = template_catalog()
    assert rows
    assert all(row.version == TEMPLATE_VERSION for row in rows)
    assert all("caption.case_name" in row.fields and "sections" in row.fields for row in rows)


def test_instantiation_preserves_outline_and_user_values() -> None:
    artifact = build_template_document(
        "motion-parenting-time-specific",
        caption={"case_name": "In re Example", "case_number": "24-0001"},
        sections=["User facts", "Requested terms"],
        signature={"name": "Owner"},
    )
    document = Document(BytesIO(artifact.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "In re Example" in table_text
    assert "24-0001" in table_text
    assert "User facts" in text
    assert "Owner" in text
    assert document.styles["Normal"].font.name == "Arial"


def test_generated_content_hash_matches_bytes_and_placeholders_remain() -> None:
    artifact = build_template_document("affidavit")
    document = Document(BytesIO(artifact.content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert artifact.content_hash == hashlib.sha256(artifact.content).hexdigest()
    assert "[Enter content for:" in text
    assert "[Enter name]" in text
